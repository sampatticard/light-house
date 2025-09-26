from fastapi import FastAPI, File, UploadFile, HTTPException, Depends, BackgroundTasks, Query
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from sqlalchemy import create_engine, Column, Integer, String, DateTime, Text, Boolean, Float, LargeBinary
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker, Session
from pydantic import BaseModel, Field
from typing import List, Optional, Dict, Any
from datetime import datetime, timedelta
from pathlib import Path
import os
import hashlib
import json
import asyncio
import aiofiles
from cryptography.fernet import Fernet
import pytesseract
from PIL import Image
import cv2
import numpy as np
from pdf2image import convert_from_bytes
import docx
import re
from enum import Enum
import logging
from contextlib import asynccontextmanager
import uvicorn


# Logging configuration
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Database configuration
DATABASE_URL = "sqlite:///./financial_agent.db"
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

# Security configuration
security = HTTPBearer()
ENCRYPTION_KEY = Fernet.generate_key()
cipher_suite = Fernet(ENCRYPTION_KEY)

# Constants
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
ALLOWED_EXTENSIONS = {'.pdf', '.jpg', '.jpeg', '.png', '.docx'}
UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

class DocumentCategory(str, Enum):
    AADHAAR = "aadhaar"
    PAN = "pan"
    BANK_STATEMENT = "bank_statement"
    SALARY_SLIP = "salary_slip"
    INCOME_TAX_RETURN = "income_tax_return"
    INSURANCE_POLICY = "insurance_policy"
    INVESTMENT_DOCUMENT = "investment_document"
    LOAN_DOCUMENT = "loan_document"
    PROPERTY_DOCUMENT = "property_document"
    OTHER = "other"

class DocumentStatus(str, Enum):
    UPLOADED = "uploaded"
    PROCESSING = "processing"
    PROCESSED = "processed"
    FAILED = "failed"
    VALIDATED = "validated"
    FLAGGED = "flagged"

# Database Models
class Document(Base):
    __tablename__ = "documents"
    
    id = Column(Integer, primary_key=True, index=True)
    filename = Column(String, nullable=False)
    original_filename = Column(String, nullable=False)
    file_path = Column(String, nullable=False)
    file_size = Column(Integer, nullable=False)
    file_hash = Column(String, unique=True, nullable=False)
    mime_type = Column(String, nullable=False)
    category = Column(String, nullable=False)
    status = Column(String, default=DocumentStatus.UPLOADED)
    upload_date = Column(DateTime, default=datetime.utcnow)
    processed_date = Column(DateTime, nullable=True)
    is_encrypted = Column(Boolean, default=True)
    ocr_text = Column(Text, nullable=True)
    extracted_data = Column(Text, nullable=True)  # JSON string
    validation_flags = Column(Text, nullable=True)  # JSON string
    confidence_score = Column(Float, default=0.0)
    language_detected = Column(String, default="en")

class ExtractedField(Base):
    __tablename__ = "extracted_fields"
    
    id = Column(Integer, primary_key=True, index=True)
    document_id = Column(Integer, nullable=False)
    field_name = Column(String, nullable=False)
    field_value = Column(String, nullable=False)
    confidence = Column(Float, default=0.0)
    field_type = Column(String, nullable=False)  # text, number, date, etc.
    extraction_method = Column(String, nullable=False)  # ocr, regex, ml
    created_date = Column(DateTime, default=datetime.utcnow)

class DocumentValidation(Base):
    __tablename__ = "document_validations"
    
    id = Column(Integer, primary_key=True, index=True)
    document_id = Column(Integer, nullable=False)
    validation_type = Column(String, nullable=False)
    is_valid = Column(Boolean, nullable=False)
    validation_message = Column(String, nullable=True)
    validation_date = Column(DateTime, default=datetime.utcnow)

# Create tables
Base.metadata.create_all(bind=engine)

# Pydantic models
class DocumentUploadResponse(BaseModel):
    id: int
    filename: str
    category: str
    status: str
    file_size: int
    upload_date: datetime
    message: str

class DocumentInfo(BaseModel):
    id: int
    filename: str
    original_filename: str
    category: str
    status: str
    file_size: int
    upload_date: datetime
    processed_date: Optional[datetime]
    confidence_score: float
    language_detected: str

class ExtractedData(BaseModel):
    document_id: int
    extracted_fields: Dict[str, Any]
    confidence_score: float
    validation_flags: List[str]

class DocumentAnalysis(BaseModel):
    document_id: int
    category: str
    extracted_data: Dict[str, Any]
    ocr_text: str
    confidence_score: float
    validation_results: List[Dict[str, Any]]
    missing_fields: List[str]
    inconsistencies: List[str]

class ValidationResult(BaseModel):
    is_valid: bool
    validation_type: str
    message: str
    confidence: float

# Dependency to get database session
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# Authentication dependency (simplified)
async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    # In production, implement proper JWT validation
    if not credentials.credentials:
        raise HTTPException(status_code=401, detail="Invalid authentication credentials")
    return {"user_id": "user123"}  # Mock user

# Utility functions
def calculate_file_hash(file_content: bytes) -> str:
    return hashlib.sha256(file_content).hexdigest()

def encrypt_file_content(content: bytes) -> bytes:
    return cipher_suite.encrypt(content)

def decrypt_file_content(encrypted_content: bytes) -> bytes:
    return cipher_suite.decrypt(encrypted_content)

def validate_file_extension(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS

def detect_document_category(filename: str, ocr_text: str = "") -> DocumentCategory:
    """Intelligent document category detection based on filename and content"""
    filename_lower = filename.lower()
    text_lower = ocr_text.lower()
    
    # Aadhaar detection
    if any(keyword in filename_lower for keyword in ['aadhaar', 'aadhar', 'uid']) or \
       any(keyword in text_lower for keyword in ['aadhaar', 'unique identification', 'uidai']):
        return DocumentCategory.AADHAAR
    
    # PAN detection
    if 'pan' in filename_lower or \
       re.search(r'[A-Z]{5}[0-9]{4}[A-Z]{1}', text_lower):
        return DocumentCategory.PAN
    
    # Bank statement detection
    if any(keyword in filename_lower for keyword in ['bank', 'statement', 'account']) or \
       any(keyword in text_lower for keyword in ['account statement', 'bank statement', 'balance']):
        return DocumentCategory.BANK_STATEMENT
    
    # Salary slip detection
    if any(keyword in filename_lower for keyword in ['salary', 'payslip', 'pay']) or \
       any(keyword in text_lower for keyword in ['salary slip', 'pay slip', 'gross salary']):
        return DocumentCategory.SALARY_SLIP
    
    return DocumentCategory.OTHER

# OCR and text extraction functions
async def extract_text_from_image(image_path: str) -> tuple[str, str]:
    """Extract text from image using OCR"""
    try:
        # Load and preprocess image
        image = cv2.imread(image_path)
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        
        # Apply image preprocessing for better OCR
        gray = cv2.medianBlur(gray, 3)
        gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
        
        # Perform OCR with both English and Hindi
        custom_config = r'--oem 3 --psm 6 -l eng+hin'
        ocr_text = pytesseract.image_to_string(gray, config=custom_config)
        
        # Detect language
        if re.search(r'[\u0900-\u097F]', ocr_text):
            language = "hi"  # Hindi detected
        else:
            language = "en"  # English
            
        return ocr_text.strip(), language
    except Exception as e:
        logger.error(f"OCR extraction error: {str(e)}")
        return "", "en"

async def extract_text_from_pdf(file_path: str) -> tuple[str, str]:
    """Extract text from PDF"""
    try:
        with open(file_path, 'rb') as file:
            file_content = file.read()
            
        # Convert PDF to images
        images = convert_from_bytes(file_content)
        extracted_text = ""
        language = "en"
        
        for image in images:
            # Convert PIL image to cv2 format
            cv_image = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2BGR)
            gray = cv2.cvtColor(cv_image, cv2.COLOR_BGR2GRAY)
            
            # OCR processing
            custom_config = r'--oem 3 --psm 6 -l eng+hin'
            page_text = pytesseract.image_to_string(gray, config=custom_config)
            extracted_text += page_text + "\n"
            
            # Check for Hindi text
            if re.search(r'[\u0900-\u097F]', page_text):
                language = "hi"
        
        return extracted_text.strip(), language
    except Exception as e:
        logger.error(f"PDF extraction error: {str(e)}")
        return "", "en"

async def extract_text_from_docx(file_path: str) -> tuple[str, str]:
    """Extract text from DOCX"""
    try:
        doc = docx.Document(file_path)
        extracted_text = ""
        
        for paragraph in doc.paragraphs:
            extracted_text += paragraph.text + "\n"
        
        # Detect language
        if re.search(r'[\u0900-\u097F]', extracted_text):
            language = "hi"
        else:
            language = "en"
            
        return extracted_text.strip(), language
    except Exception as e:
        logger.error(f"DOCX extraction error: {str(e)}")
        return "", "en"

def extract_structured_data(text: str, category: DocumentCategory) -> Dict[str, Any]:
    """Extract structured data based on document category"""
    extracted_data = {}
    
    if category == DocumentCategory.AADHAAR:
        # Aadhaar number extraction
        aadhaar_match = re.search(r'\b\d{4}\s?\d{4}\s?\d{4}\b', text)
        if aadhaar_match:
            extracted_data['aadhaar_number'] = aadhaar_match.group().replace(' ', '')
        
        # Name extraction (simplified)
        name_patterns = [
            r'Name[:\s]+([A-Za-z\s]+?)(?:\n|Date|DOB)',
            r'नाम[:\s]+([A-Za-z\s\u0900-\u097F]+?)(?:\n|जन्म|DOB)'
        ]
        for pattern in name_patterns:
            name_match = re.search(pattern, text, re.IGNORECASE)
            if name_match:
                extracted_data['name'] = name_match.group(1).strip()
                break
    
    elif category == DocumentCategory.PAN:
        # PAN number extraction
        pan_match = re.search(r'\b[A-Z]{5}[0-9]{4}[A-Z]{1}\b', text)
        if pan_match:
            extracted_data['pan_number'] = pan_match.group()
        
        # Name extraction
        name_match = re.search(r'Name[:\s]+([A-Za-z\s]+?)(?:\n|Father)', text, re.IGNORECASE)
        if name_match:
            extracted_data['name'] = name_match.group(1).strip()
    
    elif category == DocumentCategory.BANK_STATEMENT:
        # Account number extraction
        acc_patterns = [
            r'Account\s+No[:\s]+(\d+)',
            r'A/c\s+No[:\s]+(\d+)',
            r'खाता\s+संख्या[:\s]+(\d+)'
        ]
        for pattern in acc_patterns:
            acc_match = re.search(pattern, text, re.IGNORECASE)
            if acc_match:
                extracted_data['account_number'] = acc_match.group(1)
                break
        
        # Balance extraction
        balance_patterns = [
            r'Balance[:\s]+(?:Rs\.?\s*)?(\d+(?:,\d+)*(?:\.\d{2})?)',
            r'शेष[:\s]+(?:रु\.?\s*)?(\d+(?:,\d+)*(?:\.\d{2})?)'
        ]
        for pattern in balance_patterns:
            balance_match = re.search(pattern, text, re.IGNORECASE)
            if balance_match:
                extracted_data['balance'] = balance_match.group(1).replace(',', '')
                break
    
    return extracted_data

def validate_document_data(extracted_data: Dict[str, Any], category: DocumentCategory) -> List[ValidationResult]:
    """Validate extracted document data"""
    validations = []
    
    if category == DocumentCategory.AADHAAR:
        if 'aadhaar_number' in extracted_data:
            aadhaar = extracted_data['aadhaar_number']
            is_valid = len(aadhaar) == 12 and aadhaar.isdigit()
            validations.append(ValidationResult(
                is_valid=is_valid,
                validation_type="aadhaar_format",
                message="Valid Aadhaar format" if is_valid else "Invalid Aadhaar format",
                confidence=0.9 if is_valid else 0.1
            ))
    
    elif category == DocumentCategory.PAN:
        if 'pan_number' in extracted_data:
            pan = extracted_data['pan_number']
            is_valid = re.match(r'^[A-Z]{5}[0-9]{4}[A-Z]{1}$', pan) is not None
            validations.append(ValidationResult(
                is_valid=is_valid,
                validation_type="pan_format",
                message="Valid PAN format" if is_valid else "Invalid PAN format",
                confidence=0.95 if is_valid else 0.1
            ))
    
    return validations

async def process_document_async(document_id: int, file_path: str, db: Session):
    """Background task to process document"""
    try:
        document = db.query(Document).filter(Document.id == document_id).first()
        if not document:
            return
        
        # Update status to processing
        document.status = DocumentStatus.PROCESSING
        db.commit()
        
        # Extract text based on file type
        file_extension = Path(file_path).suffix.lower()
        
        if file_extension in ['.jpg', '.jpeg', '.png']:
            ocr_text, language = await extract_text_from_image(file_path)
        elif file_extension == '.pdf':
            ocr_text, language = await extract_text_from_pdf(file_path)
        elif file_extension == '.docx':
            ocr_text, language = await extract_text_from_docx(file_path)
        else:
            ocr_text, language = "", "en"
        
        # Detect category if not manually set
        if document.category == DocumentCategory.OTHER:
            detected_category = detect_document_category(document.original_filename, ocr_text)
            document.category = detected_category
        
        # Extract structured data
        extracted_data = extract_structured_data(ocr_text, DocumentCategory(document.category))
        
        # Calculate confidence score
        confidence_score = min(1.0, len(ocr_text) / 1000)  # Simplified confidence calculation
        
        # Validate document
        validation_results = validate_document_data(extracted_data, DocumentCategory(document.category))
        
        # Update document record
        document.ocr_text = ocr_text
        document.extracted_data = json.dumps(extracted_data)
        document.confidence_score = confidence_score
        document.language_detected = language
        document.processed_date = datetime.utcnow()
        document.status = DocumentStatus.PROCESSED
        
        # Store extracted fields
        for field_name, field_value in extracted_data.items():
            extracted_field = ExtractedField(
                document_id=document_id,
                field_name=field_name,
                field_value=str(field_value),
                confidence=confidence_score,
                field_type="text",
                extraction_method="ocr"
            )
            db.add(extracted_field)
        
        # Store validation results
        for validation in validation_results:
            doc_validation = DocumentValidation(
                document_id=document_id,
                validation_type=validation.validation_type,
                is_valid=validation.is_valid,
                validation_message=validation.message
            )
            db.add(doc_validation)
        
        db.commit()
        logger.info(f"Document {document_id} processed successfully")
        
    except Exception as e:
        logger.error(f"Error processing document {document_id}: {str(e)}")
        document.status = DocumentStatus.FAILED
        db.commit()

# FastAPI app initialization
@asynccontextmanager
async def lifespan(app: FastAPI):
    logger.info("Financial Agent DMS Backend starting up...")
    yield
    logger.info("Financial Agent DMS Backend shutting down...")

app = FastAPI(
    title="Financial Agent - Document Management System",
    description="AI-powered document management system for financial documents",
    version="1.0.0",
    docs_url="/docs",           # Swagger UI
    redoc_url="/redoc",         # ReDoc UI
    openapi_url="/openapi.json", # OpenAPI schema
    lifespan=lifespan
)

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Configure appropriately for production
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# API Endpoints

@app.get("/")
async def root():
    return {
        "message": "Financial Agent Document Management System API",
        "version": "1.0.0",
        "swagger_docs": "/docs",
        "redoc_docs": "/redoc",
        "openapi_schema": "/openapi.json"
    }

@app.post("/documents/upload", response_model=DocumentUploadResponse)
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    category: Optional[DocumentCategory] = Query(DocumentCategory.OTHER),
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Upload and process a document"""
    
    # Validate file
    if not validate_file_extension(file.filename):
        raise HTTPException(status_code=400, detail="Unsupported file format")
    
    if file.size > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File size exceeds 10MB limit")
    
    try:
        # Read file content
        file_content = await file.read()
        file_hash = calculate_file_hash(file_content)
        
        # Check for duplicate
        existing_doc = db.query(Document).filter(Document.file_hash == file_hash).first()
        if existing_doc:
            raise HTTPException(status_code=400, detail="Document already exists")
        
        # Generate unique filename
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        file_extension = Path(file.filename).suffix
        unique_filename = f"{timestamp}_{file_hash[:8]}{file_extension}"
        file_path = UPLOAD_DIR / unique_filename
        
        # Encrypt and save file
        encrypted_content = encrypt_file_content(file_content)
        async with aiofiles.open(file_path, 'wb') as f:
            await f.write(encrypted_content)
        
        # Create document record
        document = Document(
            filename=unique_filename,
            original_filename=file.filename,
            file_path=str(file_path),
            file_size=file.size,
            file_hash=file_hash,
            mime_type=file.content_type,
            category=category,
            status=DocumentStatus.UPLOADED,
            is_encrypted=True
        )
        
        db.add(document)
        db.commit()
        db.refresh(document)
        
        # Start background processing
        background_tasks.add_task(process_document_async, document.id, str(file_path), db)
        
        return DocumentUploadResponse(
            id=document.id,
            filename=document.filename,
            category=document.category,
            status=document.status,
            file_size=document.file_size,
            upload_date=document.upload_date,
            message="Document uploaded successfully. Processing started."
        )
        
    except Exception as e:
        logger.error(f"Upload error: {str(e)}")
        raise HTTPException(status_code=500, detail="File upload failed")

@app.get("/documents", response_model=List[DocumentInfo])
async def list_documents(
    category: Optional[DocumentCategory] = None,
    status: Optional[DocumentStatus] = None,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """List all documents with optional filtering"""
    
    query = db.query(Document)
    
    if category:
        query = query.filter(Document.category == category)
    if status:
        query = query.filter(Document.status == status)
    
    documents = query.order_by(Document.upload_date.desc()).all()
    
    return [
        DocumentInfo(
            id=doc.id,
            filename=doc.filename,
            original_filename=doc.original_filename,
            category=doc.category,
            status=doc.status,
            file_size=doc.file_size,
            upload_date=doc.upload_date,
            processed_date=doc.processed_date,
            confidence_score=doc.confidence_score,
            language_detected=doc.language_detected
        )
        for doc in documents
    ]

@app.get("/documents/{document_id}/analysis", response_model=DocumentAnalysis)
async def get_document_analysis(
    document_id: int,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get detailed analysis of a processed document"""
    
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    if document.status != DocumentStatus.PROCESSED:
        raise HTTPException(status_code=400, detail="Document not yet processed")
    
    # Get extracted data
    extracted_data = json.loads(document.extracted_data) if document.extracted_data else {}
    
    # Get validation results
    validations = db.query(DocumentValidation).filter(
        DocumentValidation.document_id == document_id
    ).all()
    
    validation_results = [
        {
            "validation_type": v.validation_type,
            "is_valid": v.is_valid,
            "message": v.validation_message,
            "date": v.validation_date
        }
        for v in validations
    ]
    
    # Identify missing fields based on category
    missing_fields = []
    inconsistencies = []
    
    if document.category == DocumentCategory.AADHAAR:
        required_fields = ['aadhaar_number', 'name']
        missing_fields = [field for field in required_fields if field not in extracted_data]
    elif document.category == DocumentCategory.PAN:
        required_fields = ['pan_number', 'name']
        missing_fields = [field for field in required_fields if field not in extracted_data]
    
    return DocumentAnalysis(
        document_id=document_id,
        category=document.category,
        extracted_data=extracted_data,
        ocr_text=document.ocr_text or "",
        confidence_score=document.confidence_score,
        validation_results=validation_results,
        missing_fields=missing_fields,
        inconsistencies=inconsistencies
    )

@app.get("/documents/{document_id}/extracted-data", response_model=ExtractedData)
async def get_extracted_data(
    document_id: int,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get extracted data for a specific document"""
    
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    extracted_data = json.loads(document.extracted_data) if document.extracted_data else {}
    
    # Get validation flags
    validations = db.query(DocumentValidation).filter(
        DocumentValidation.document_id == document_id
    ).all()
    
    validation_flags = [v.validation_type for v in validations if not v.is_valid]
    
    return ExtractedData(
        document_id=document_id,
        extracted_fields=extracted_data,
        confidence_score=document.confidence_score,
        validation_flags=validation_flags
    )

@app.delete("/documents/{document_id}")
async def delete_document(
    document_id: int,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Delete a document and its associated data"""
    
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    try:
        # Delete physical file
        if os.path.exists(document.file_path):
            os.remove(document.file_path)
        
        # Delete related records
        db.query(ExtractedField).filter(ExtractedField.document_id == document_id).delete()
        db.query(DocumentValidation).filter(DocumentValidation.document_id == document_id).delete()
        
        # Delete document record
        db.delete(document)
        db.commit()
        
        return {"message": "Document deleted successfully"}
        
    except Exception as e:
        logger.error(f"Delete error: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to delete document")

@app.post("/documents/{document_id}/reprocess")
async def reprocess_document(
    document_id: int,
    background_tasks: BackgroundTasks,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Reprocess a document"""
    
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    # Reset document status
    document.status = DocumentStatus.UPLOADED
    document.processed_date = None
    document.ocr_text = None
    document.extracted_data = None
    db.commit()
    
    # Clear existing extracted fields and validations
    db.query(ExtractedField).filter(ExtractedField.document_id == document_id).delete()
    db.query(DocumentValidation).filter(DocumentValidation.document_id == document_id).delete()
    db.commit()
    
    # Start reprocessing
    background_tasks.add_task(process_document_async, document_id, document.file_path, db)
    
    return {"message": "Document reprocessing started"}

@app.get("/documents/stats")
async def get_document_stats(
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get document statistics"""
    
    total_documents = db.query(Document).count()
    processed_documents = db.query(Document).filter(Document.status == DocumentStatus.PROCESSED).count()
    failed_documents = db.query(Document).filter(Document.status == DocumentStatus.FAILED).count()
    
    # Category breakdown
    category_stats = {}
    for category in DocumentCategory:
        count = db.query(Document).filter(Document.category == category).count()
        category_stats[category.value] = count
    
    return {
        "total_documents": total_documents,
        "processed_documents": processed_documents,
        "failed_documents": failed_documents,
        "processing_documents": total_documents - processed_documents - failed_documents,
        "category_breakdown": category_stats
    }

@app.get("/documents/cross-reference")
async def cross_reference_documents(
    field_name: str,
    current_user: dict = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Cross-reference data across multiple documents"""
    
    extracted_fields = db.query(ExtractedField).filter(
        ExtractedField.field_name == field_name
    ).all()
    
    field_values = {}
    for field in extracted_fields:
        document = db.query(Document).filter(Document.id == field.document_id).first()
        if field.field_value not in field_values:
            field_values[field.field_value] = []
        
        field_values[field.field_value].append({
            "document_id": field.document_id,
            "document_name": document.original_filename,
            "category": document.category,
            "confidence": field.confidence
        })
    
    # Identify inconsistencies
    inconsistencies = []
    if len(field_values) > 1:
        inconsistencies.append({
            "field_name": field_name,
            "conflicting_values": list(field_values.keys()),
            "message": f"Inconsistent {field_name} values found across documents"
        })
    
    return {
        "field_name": field_name,
        "unique_values": len(field_values),
        "field_occurrences": field_values,
        "inconsistencies": inconsistencies
    }

@app.get("/health")
async def health_check():
    """Health check endpoint"""
    return {
        "status": "healthy",
        "timestamp": datetime.utcnow(),
        "version": "1.0.0"
    }

async def enhanced_process_document(document_id: int, file_path: str, db: Session):
    """Enhanced document processing with category-specific extraction"""
    try:
        document = db.query(Document).filter(Document.id == document_id).first()
        if not document:
            return
        
        document.status = DocumentStatus.PROCESSING
        db.commit()
        
        # Decrypt file for processing
        with open(file_path, 'rb') as f:
            encrypted_content = f.read()
        decrypted_content = decrypt_file_content(encrypted_content)
        
        # Save decrypted file temporarily
        temp_path = file_path + '.temp'
        with open(temp_path, 'wb') as f:
            f.write(decrypted_content)
        
        try:
            file_extension = Path(file_path).suffix.lower()
            
            # Enhanced OCR with preprocessing
            if file_extension in ['.jpg', '.jpeg', '.png']:
                preprocessed_path = DocumentProcessor.preprocess_image_for_ocr(temp_path)
                ocr_text, language = await extract_text_from_image(preprocessed_path)
                os.remove(preprocessed_path)  # Clean up
            elif file_extension == '.pdf':
                ocr_text, language = await extract_text_from_pdf(temp_path)
            elif file_extension == '.docx':
                ocr_text, language = await extract_text_from_docx(temp_path)
            else:
                ocr_text, language = "", "en"
            
            # Category-specific extraction
            if document.category == DocumentCategory.AADHAAR:
                extracted_data = DocumentProcessor.extract_aadhaar_details(ocr_text)
            elif document.category == DocumentCategory.PAN:
                extracted_data = DocumentProcessor.extract_pan_details(ocr_text)
            elif document.category == DocumentCategory.BANK_STATEMENT:
                extracted_data = DocumentProcessor.extract_bank_statement_details(ocr_text)
            else:
                extracted_data = extract_structured_data(ocr_text, DocumentCategory(document.category))
            
            # Enhanced validation
            validation_results = []
            if document.category == DocumentCategory.AADHAAR and 'aadhaar_number' in extracted_data:
                is_valid = DocumentValidator.validate_aadhaar_checksum(extracted_data['aadhaar_number'])
                validation_results.append(ValidationResult(
                    is_valid=is_valid,
                    validation_type="aadhaar_checksum",
                    message="Valid Aadhaar checksum" if is_valid else "Invalid Aadhaar checksum",
                    confidence=0.95
                ))
            
            elif document.category == DocumentCategory.PAN and 'pan_number' in extracted_data:
                is_valid = DocumentValidator.validate_pan_checksum(extracted_data['pan_number'])
                validation_results.append(ValidationResult(
                    is_valid=is_valid,
                    validation_type="pan_format",
                    message="Valid PAN format" if is_valid else "Invalid PAN format",
                    confidence=0.95
                ))
            
            # Calculate confidence based on extraction success
            confidence_score = len(extracted_data) / 5  # Normalize by expected fields
            confidence_score = min(1.0, confidence_score)
            
            # Update document
            document.ocr_text = ocr_text
            document.extracted_data = json.dumps(extracted_data)
            document.confidence_score = confidence_score
            document.language_detected = language
            document.processed_date = datetime.utcnow()
            document.status = DocumentStatus.PROCESSED
            
            # Store results
            for field_name, field_value in extracted_data.items():
                extracted_field = ExtractedField(
                    document_id=document_id,
                    field_name=field_name,
                    field_value=str(field_value),
                    confidence=confidence_score,
                    field_type="text",
                    extraction_method="enhanced_ocr"
                )
                db.add(extracted_field)
            
            for validation in validation_results:
                doc_validation = DocumentValidation(
                    document_id=document_id,
                    validation_type=validation.validation_type,
                    is_valid=validation.is_valid,
                    validation_message=validation.message
                )
                db.add(doc_validation)
            
            db.commit()
            logger.info(f"Document {document_id} processed successfully with enhanced extraction")
            
        finally:
            # Clean up temporary file
            if os.path.exists(temp_path):
                os.remove(temp_path)
        
    except Exception as e:
        logger.error(f"Enhanced processing error for document {document_id}: {str(e)}")
        document.status = DocumentStatus.FAILED
        db.commit()

class DocumentProcessor:
    """Advanced document processing utilities"""
    
    @staticmethod
    def preprocess_image_for_ocr(image_path: str) -> str:
        """Advanced image preprocessing for better OCR results"""
        image = cv2.imread(image_path)
        
        # Convert to grayscale
        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        
        # Noise reduction
        denoised = cv2.fastNlMeansDenoising(gray)
        
        # Deskewing
        coords = np.column_stack(np.where(denoised > 0))
        angle = cv2.minAreaRect(coords)[-1]
        if angle < -45:
            angle = -(90 + angle)
        else:
            angle = -angle
        
        (h, w) = denoised.shape[:2]
        center = (w // 2, h // 2)
        M = cv2.getRotationMatrix2D(center, angle, 1.0)
        rotated = cv2.warpAffine(denoised, M, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)
        
        # Contrast enhancement
        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
        enhanced = clahe.apply(rotated)
        
        # Save preprocessed image
        preprocessed_path = image_path.replace('.', '_preprocessed.')
        cv.imwrite(preprocessed_path, enhanced)
        
        return preprocessed_path

    @staticmethod
    def extract_aadhaar_details(text: str) -> Dict[str, Any]:
        """Extract comprehensive Aadhaar card details"""
        details = {}
        
        # Aadhaar number
        aadhaar_pattern = r'\b\d{4}\s?\d{4}\s?\d{4}\b'
        aadhaar_match = re.search(aadhaar_pattern, text)
        if aadhaar_match:
            details['aadhaar_number'] = aadhaar_match.group().replace(' ', '')
        
        # Name (English)
        name_patterns = [
            r'Name[:\s]*([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)',
            r'([A-Z][A-Z\s]{2,}?)(?:\s*(?:S/O|D/O|W/O))',
        ]
        
        for pattern in name_patterns:
            name_match = re.search(pattern, text)
            if name_match:
                details['name'] = name_match.group(1).strip()
                break
        
        # Father's/Spouse name
        father_patterns = [
            r'(?:S/O|Son of|Father)[:\s]*([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)',
            r'(?:D/O|Daughter of)[:\s]*([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)',
            r'(?:W/O|Wife of)[:\s]*([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)',
        ]
        
        for pattern in father_patterns:
            father_match = re.search(pattern, text, re.IGNORECASE)
            if father_match:
                details['father_name'] = father_match.group(1).strip()
                break
        
        # Date of Birth
        dob_patterns = [
            r'DOB[:\s]*(\d{2}/\d{2}/\d{4})',
            r'Date of Birth[:\s]*(\d{2}/\d{2}/\d{4})',
            r'जन्म तिथि[:\s]*(\d{2}/\d{2}/\d{4})',
        ]
        
        for pattern in dob_patterns:
            dob_match = re.search(pattern, text, re.IGNORECASE)
            if dob_match:
                details['date_of_birth'] = dob_match.group(1)
                break
        
        # Gender
        if re.search(r'\b(?:MALE|पुरुष)\b', text, re.IGNORECASE):
            details['gender'] = 'Male'
        elif re.search(r'\b(?:FEMALE|महिला)\b', text, re.IGNORECASE):
            details['gender'] = 'Female'
        
        # Address
        address_patterns = [
            r'Address[:\s]*([^\n]{20,})',
            r'पता[:\s]*([^\n]{20,})',
        ]
        
        for pattern in address_patterns:
            address_match = re.search(pattern, text, re.IGNORECASE)
            if address_match:
                details['address'] = address_match.group(1).strip()
                break
        
        return details

    @staticmethod
    def extract_pan_details(text: str) -> Dict[str, Any]:
        """Extract comprehensive PAN card details"""
        details = {}
        
        # PAN number
        pan_pattern = r'\b[A-Z]{5}[0-9]{4}[A-Z]{1}\b'
        pan_match = re.search(pan_pattern, text)
        if pan_match:
            details['pan_number'] = pan_match.group()
        
        # Name
        name_patterns = [
            r'Name[:\s]*([A-Z][A-Z\s]{2,}?)(?:\s*Father)',
            r'([A-Z][A-Z\s]{10,}?)(?:\s*Father)',
        ]
        
        for pattern in name_patterns:
            name_match = re.search(pattern, text)
            if name_match:
                details['name'] = name_match.group(1).strip()
                break
        
        # Father's name
        father_patterns = [
            r"Father's Name[:\s]*([A-Z][A-Z\s]{2,}?)(?:\s*Date|$)",
            r"Father[:\s]*([A-Z][A-Z\s]{2,}?)(?:\s*Date|$)",
        ]
        
        for pattern in father_patterns:
            father_match = re.search(pattern, text)
            if father_match:
                details['father_name'] = father_match.group(1).strip()
                break
        
        # Date of Birth
        dob_patterns = [
            r'Date of Birth[:\s]*(\d{2}/\d{2}/\d{4})',
            r'DOB[:\s]*(\d{2}/\d{2}/\d{4})',
        ]
        
        for pattern in dob_patterns:
            dob_match = re.search(pattern, text)
            if dob_match:
                details['date_of_birth'] = dob_match.group(1)
                break
        
        return details

    @staticmethod
    def extract_bank_statement_details(text: str) -> Dict[str, Any]:
        """Extract bank statement details"""
        details = {}
        
        # Account number
        account_patterns = [
            r'Account\s+No[:\s]*(\d+)',
            r'A/c\s+No[:\s]*(\d+)',
            r'Account\s+Number[:\s]*(\d+)',
        ]
        
        for pattern in account_patterns:
            acc_match = re.search(pattern, text, re.IGNORECASE)
            if acc_match:
                details['account_number'] = acc_match.group(1).strip()
                break
        
        # IFSC code
        ifsc_pattern = r'\b[A-Z]{4}0[A-Z0-9]{6}\b'
        ifsc_match = re.search(ifsc_pattern, text)
        if ifsc_match:
            details['ifsc_code'] = ifsc_match.group()
        
        # Account holder name
        name_patterns = [
            r'Account\s+Holder[:\s]*([A-Z][A-Za-z\s]{2,}?)(?:\s*Account|\n)',
            r'Name[:\s]*([A-Z][A-Za-z\s]{2,}?)(?:\s*Account|\n)',
        ]
        
        for pattern in name_patterns:
            name_match = re.search(pattern, text, re.IGNORECASE)
            if name_match:
                details['account_holder_name'] = name_match.group(1).strip()
                break
        
        # Balance
        balance_patterns = [
            r'Balance[:\s]*(?:Rs\.?\s*)?(\d+(?:,\d+)*(?:\.\d{2})?)',
            r'Available\s+Balance[:\s]*(?:Rs\.?\s*)?(\d+(?:,\d+)*(?:\.\d{2})?)',
            r'Closing\s+Balance[:\s]*(?:Rs\.?\s*)?(\d+(?:,\d+)*(?:\.\d{2})?)',
        ]
        
        for pattern in balance_patterns:
            balance_match = re.search(pattern, text, re.IGNORECASE)
            if balance_match:
                details['balance'] = balance_match.group(1).replace(',', '')
                break
        
        # Statement period
        period_patterns = [
            r'Statement\s+Period[:\s]*(\d{2}/\d{2}/\d{4})\s*to\s*(\d{2}/\d{2}/\d{4})',
            r'From[:\s]*(\d{2}/\d{2}/\d{4})\s*To[:\s]*(\d{2}/\d{2}/\d{4})',
        ]
        
        for pattern in period_patterns:
            period_match = re.search(pattern, text, re.IGNORECASE)
            if period_match:
                details['statement_from'] = period_match.group(1)
                details['statement_to'] = period_match.group(2)
                break
        
        return details


class DocumentValidator:
    """Advanced document validation utilities"""
    
    @staticmethod
    def validate_aadhaar_checksum(aadhaar_number: str) -> bool:
        """Validate Aadhaar number using Verhoeff algorithm"""
        if len(aadhaar_number) != 12 or not aadhaar_number.isdigit():
            return False
        
        # Verhoeff algorithm implementation
        multiplication_table = [
            [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
            [1, 2, 3, 4, 0, 6, 7, 8, 9, 5],
            [2, 3, 4, 0, 1, 7, 8, 9, 5, 6],
            [3, 4, 0, 1, 2, 8, 9, 5, 6, 7],
            [4, 0, 1, 2, 3, 9, 5, 6, 7, 8],
            [5, 9, 8, 7, 6, 0, 4, 3, 2, 1],
            [6, 5, 9, 8, 7, 1, 0, 4, 3, 2],
            [7, 6, 5, 9, 8, 2, 1, 0, 4, 3],
            [8, 7, 6, 5, 9, 3, 2, 1, 0, 4],
            [9, 8, 7, 6, 5, 4, 3, 2, 1, 0]
        ]
        
        permutation_table = [
            [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
            [1, 5, 7, 6, 2, 8, 3, 0, 9, 4],
            [5, 8, 0, 3, 7, 9, 6, 1, 4, 2],
            [8, 9, 1, 6, 0, 4, 3, 5, 2, 7],
            [9, 4, 5, 3, 1, 2, 6, 8, 7, 0],
            [4, 2, 8, 6, 5, 7, 3, 9, 0, 1],
            [2, 7, 9, 3, 8, 0, 6, 4, 1, 5],
            [7, 0, 4, 6, 9, 1, 3, 2, 5, 8]
        ]
        
        check = 0
        for i, digit in enumerate(reversed(aadhaar_number)):
            check = multiplication_table[check][permutation_table[i % 8][int(digit)]]
        
        return check == 0
    
    @staticmethod
    def validate_pan_checksum(pan_number: str) -> bool:
        """Validate PAN number format and structure"""
        if len(pan_number) != 10:
            return False
        
        # Check format: AAAPL1234C
        pattern = r'^[A-Z]{5}[0-9]{4}[A-Z]{1}'
        if not re.match(pattern, pan_number):
            return False
        
        # Fourth character should be 'P' for individual PAN
        if pan_number[3] not in ['P', 'F', 'A', 'T', 'B', 'C', 'G', 'H', 'L', 'J']:
            return False
        
        return True
    
    @staticmethod
    def validate_ifsc_code(ifsc_code: str) -> bool:
        """Validate IFSC code format"""
        if len(ifsc_code) != 11:
            return False
        
        # Format: ABCD0123456 (first 4 alpha, 5th is 0, last 6 alphanumeric)
        pattern = r'^[A-Z]{4}0[A-Z0-9]{6}'
        return bool(re.match(pattern, ifsc_code))

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)